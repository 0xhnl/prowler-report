"""
Prowler -> preliminary cloud pentest report.

Usage:
    python3 final.py -f raw-report/ -o out-report/
    python3 final.py -f raw-report/ -o out-report/ \
        --project "Yoma Elevator" --start "01 June 2026" --end "04 June 2026"

Output (in the output folder):
    rpt.<code>-<project>-cloud-pentest-prelim.<date>.xlsx    all findings (one row per Prowler result)
    rpt.<code>-<project>-cloud-pentest-prelim.<date>.docx    report built from template.docx
    finding-count.xlsx                                       finding counts by severity/service/region

Source: the native Prowler CSV is preferred because it is the only output that
carries remediation code (CloudFormation, Terraform, CLI, console steps). The
HTML report is used as a fallback when no CSV exists for a scan.

The report is built from template.docx (next to this script). The template
contains {{PLACEHOLDERS}} for the project details, a {{FINDINGS}} marker where
the findings are inserted, and three native charts that are refreshed with the
scan numbers.
"""
import argparse
import io
import re
import sys
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree


DEFAULT_TEMPLATE = Path(__file__).resolve().parent / "template.docx"

HTML_COLUMNS_TO_EXTRACT = [0, 1, 2, 3, 5, 6, 8, 9, 10]
CSV_HEADER = [
    "Status",
    "Severity",
    "Service Name",
    "Region",
    "Check Title",
    "Resource ID",
    "Status Extended",
    "Risk",
    "Recommendation",
    "Reference URL",
    "Remediation Steps",
    "CLI",
    "Terraform",
    "CloudFormation",
]
# Prowler native CSV column -> our column
PROWLER_CSV_MAP = {
    "STATUS": "Status",
    "SEVERITY": "Severity",
    "SERVICE_NAME": "Service Name",
    "REGION": "Region",
    "CHECK_TITLE": "Check Title",
    "RESOURCE_UID": "Resource ID",
    "STATUS_EXTENDED": "Status Extended",
    "RISK": "Risk",
    "REMEDIATION_RECOMMENDATION_TEXT": "Recommendation",
    "REMEDIATION_RECOMMENDATION_URL": "Reference URL",
    "REMEDIATION_CODE_OTHER": "Remediation Steps",
    "REMEDIATION_CODE_CLI": "CLI",
    "REMEDIATION_CODE_TERRAFORM": "Terraform",
    "REMEDIATION_CODE_NATIVEIAC": "CloudFormation",
}
HTML_COLUMNS = CSV_HEADER[:9]
SEVERITIES = ["critical", "high", "medium", "low"]
CODE_COLUMNS = ["Remediation Steps", "CLI", "Terraform", "CloudFormation"]

# Display names for AWS services in the narrative text
SERVICE_NAMES = {
    "accessanalyzer": "IAM Access Analyzer", "acm": "ACM", "apigateway": "API Gateway",
    "awslambda": "Lambda", "backup": "AWS Backup", "bedrock": "Bedrock",
    "cloudformation": "CloudFormation", "cloudfront": "CloudFront", "cloudtrail": "CloudTrail",
    "cloudwatch": "CloudWatch", "config": "AWS Config", "dynamodb": "DynamoDB", "ec2": "EC2",
    "ecr": "ECR", "ecs": "ECS", "efs": "EFS", "eks": "EKS", "elb": "ELB", "elbv2": "ELBv2",
    "guardduty": "GuardDuty", "iam": "IAM", "inspector2": "Inspector", "kms": "KMS",
    "networkfirewall": "Network Firewall", "organizations": "Organizations", "rds": "RDS",
    "route53": "Route 53", "s3": "S3", "secretsmanager": "Secrets Manager",
    "securityhub": "Security Hub", "sns": "SNS", "sqs": "SQS", "ssm": "Systems Manager",
    "vpc": "VPC", "wafv2": "WAF", "account": "Account",
}

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NSMAP_C = {"c": C_NS}


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def clean_text(text):
    """Remove HTML artifacts and normalize whitespace."""
    text = text.replace("<wbr />", "")
    soup = BeautifulSoup(text, "html.parser")
    text = soup.get_text(separator=" ", strip=True)
    return re.sub(r"\s+", " ", text).strip()


def parse_prowler_html(input_file):
    """Parse the Prowler HTML report into a DataFrame."""
    html_content = Path(input_file).read_text(encoding="utf-8")
    soup = BeautifulSoup(html_content, "html.parser")

    findings_table = soup.find("table", id="findingsTable")
    if not findings_table:
        raise ValueError("Could not find the main findings table (id='findingsTable') in the HTML.")

    tbody = findings_table.find("tbody")
    if not tbody:
        raise ValueError("Could not find the findings table body in the HTML.")

    findings = []
    for row in tbody.find_all("tr"):
        cells = row.find_all("td")
        if len(cells) < 11:
            continue
        values = [clean_text(cells[i].decode_contents()) for i in HTML_COLUMNS_TO_EXTRACT]
        link = cells[10].find("a", href=True)
        findings.append(values + [link["href"] if link else ""])

    # The HTML report has no remediation code; those columns stay empty.
    df = pd.DataFrame(findings, columns=HTML_COLUMNS + ["Reference URL"])
    return df.reindex(columns=CSV_HEADER, fill_value=""), {}


def parse_prowler_csv(input_file):
    """Parse a native Prowler CSV (semicolon separated) into the same DataFrame layout."""
    raw = pd.read_csv(input_file, sep=";", dtype=str, keep_default_na=False)
    missing = [c for c in PROWLER_CSV_MAP if c not in raw.columns]
    if missing:
        raise ValueError(f"Not a Prowler CSV, missing columns: {', '.join(missing)}")

    meta = {}
    if len(raw):
        first = raw.iloc[0]
        meta["account"] = first.get("ACCOUNT_UID", "")
        meta["prowler_version"] = first.get("PROWLER_VERSION", "")
        try:
            meta["scan_date"] = datetime.fromisoformat(first.get("TIMESTAMP", "")).date()
        except ValueError:
            pass
    return raw[list(PROWLER_CSV_MAP)].rename(columns=PROWLER_CSV_MAP), meta


# --------------------------------------------------------------------------- #
# Finding counts XLSX
# --------------------------------------------------------------------------- #
def status_pivot(df, index):
    pivot = df.pivot_table(index=index, columns="Status", aggfunc="size", fill_value=0)
    pivot["Total"] = pivot.sum(axis=1)
    sort_col = "FAIL" if "FAIL" in pivot.columns else "Total"
    return pivot.sort_values(sort_col, ascending=False).reset_index()


def severity_counts(fails):
    """Unique failed checks (= findings in the report) and FAIL rows per severity."""
    sev = fails["Severity"].str.lower()
    rows = sev.value_counts()
    checks = fails.groupby(sev)["Check Title"].nunique()
    order = SEVERITIES + sorted(set(sev) - set(SEVERITIES))
    return pd.DataFrame({
        "Severity": [s.title() for s in order],
        "Findings": [int(checks.get(s, 0)) for s in order],
        "Affected Resources": [int(rows.get(s, 0)) for s in order],
    })


def build_count_xlsx(df, output_xlsx):
    fails = df[df["Status"] == "FAIL"]
    by_service = status_pivot(df, "Service Name")
    by_region = status_pivot(df, "Region")
    by_severity = severity_counts(fails)

    if fails.empty:
        fail_service_region = pd.DataFrame(columns=["Service Name", "Total"])
    else:
        fail_service_region = fails.pivot_table(
            index="Service Name", columns="Region", aggfunc="size", fill_value=0
        )
        fail_service_region["Total"] = fail_service_region.sum(axis=1)
        fail_service_region = fail_service_region.sort_values("Total", ascending=False).reset_index()

    sheets = [
        ("FAIL By Severity", by_severity),
        ("By Service", by_service),
        ("By Region", by_region),
        ("FAIL Service x Region", fail_service_region),
    ]
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        for sheet_name, frame in sheets:
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            for col_idx, col in enumerate(frame.columns, start=1):
                values = [len(str(v)) for v in frame[col].astype(str)]
                width = max([len(str(col)), *values]) + 2
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 40)

    return by_service, by_region, by_severity


def build_findings_xlsx(df, output_xlsx):
    """All findings, FAIL first, then by severity and check title."""
    sev_rank = {s: i for i, s in enumerate(SEVERITIES)}
    findings = df.assign(
        _status=(df["Status"] != "FAIL").astype(int),
        _sev=df["Severity"].str.lower().map(sev_rank).fillna(len(SEVERITIES)),
    ).sort_values(["_status", "_sev", "Check Title"], kind="stable").drop(columns=["_status", "_sev"])

    widths = {"Status": 10, "Severity": 12, "Service Name": 16, "Region": 16, "Check Title": 45,
              "Resource ID": 50, "Reference URL": 40}
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        findings.to_excel(writer, sheet_name="Findings", index=False)
        ws = writer.sheets["Findings"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wrap = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        for col_idx, col in enumerate(findings.columns, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = widths.get(col, 60)
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                cell.alignment = wrap
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
    return len(findings)


# --------------------------------------------------------------------------- #
# DOCX helpers
# --------------------------------------------------------------------------- #
def strip_markdown(text):
    """Drop markdown emphasis/inline-code markers used in Prowler text fields."""
    return re.sub(r"\*\*|`", "", text).strip()


def strip_code_fence(text):
    """Remove ```lang ... ``` fences around Prowler remediation code."""
    text = text.strip()
    text = re.sub(r"^```[\w-]*\s*\n?", "", text)
    text = re.sub(r"\n?```\s*$", "", text)
    return text.strip()


def split_code_blocks(text):
    """Split text into (is_code, text) chunks around ```fenced``` blocks."""
    chunks = []
    for i, part in enumerate(re.split(r"```[\w-]*[ \t]*\n?", text)):
        if part.strip():
            chunks.append((i % 2 == 1, part.strip("\n")))
    return chunks


def add_sub_bullet_numbering(doc):
    """The template's List Bullet has a single level; add a hollow "o" bullet indented one
    level deeper for the "- item" lines Prowler uses inside Risk/Recommendation text."""
    numbering = doc.part.numbering_part.element
    abstract_id = max(int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))) + 1
    nums = numbering.findall(qn("w:num"))
    num_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
    nums[0].addprevious(parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New" w:hint="default"/></w:rPr>'
        "</w:lvl></w:abstractNum>"
    ))
    nums[-1].addnext(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
    ))
    return num_id


def add_bullet_text(marker, text, sub_num_id):
    """Plain lines -> main bullets; "- item" lines after a main bullet -> secondary bullets."""
    first = True
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        item = re.match(r"^[-*•]\s+(.*)", line)
        if item and not first:
            p = marker.insert_paragraph_before(style="List Bullet")
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = sub_num_id
            p.add_run(strip_markdown(item.group(1)))
        else:
            marker.insert_paragraph_before(style="List Bullet").add_run(strip_markdown(item.group(1) if item else line))
        first = False


def unique_values(group, column):
    return [v for v in group[column].dropna().astype(str).str.strip().unique() if v]


def replace_in_paragraph(p_el, old, new):
    """Replace text in a w:p even when Word split it across several runs."""
    # Only this paragraph's own text, not that of text boxes nested inside it
    ts = [t for t in p_el.iter(qn("w:t")) if next(t.iterancestors(qn("w:p"))) is p_el]
    while True:
        text = "".join(t.text or "" for t in ts)
        start = text.find(old)
        if start < 0:
            return
        end = start + len(old)
        pos, placed = 0, False
        for t in ts:
            t_text = t.text or ""
            t_start, t_end = pos, pos + len(t_text)
            pos = t_end
            if t_end <= start or t_start >= end:
                continue
            before = t_text[: max(0, start - t_start)]
            after = t_text[end - t_start:] if end < t_end else ""
            t.text = before + ("" if placed else new) + after
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            placed = True


def replace_everywhere(doc, mapping):
    """Replace placeholders in the body (incl. text boxes), headers and footers."""
    roots = [doc.element.body]
    for part in doc.part.package.parts:
        if re.search(r"/word/(header|footer)\d*\.xml$", str(part.partname)):
            roots.append(part.element)
    for root in roots:
        for p_el in root.iter(qn("w:p")):
            if "{{" not in "".join(t.text or "" for t in p_el.iter(qn("w:t"))):
                continue
            for old, new in mapping.items():
                replace_in_paragraph(p_el, old, new)


def collapse_repeated_words(doc, words=("cloud", "aws")):
    """Template text around {{PROJECT}} already says "cloud", so a project named
    "EM Cloud" would read "EM Cloud cloud environment". Keep the first word only."""
    pattern = re.compile(rf"\b({'|'.join(words)})\s+\1\b", re.IGNORECASE)
    roots = [doc.element.body]
    for part in doc.part.package.parts:
        if re.search(r"/word/(header|footer)\d*\.xml$", str(part.partname)):
            roots.append(part.element)
    for root in roots:
        for p_el in root.iter(qn("w:p")):
            text = "".join(t.text or "" for t in p_el.iter(qn("w:t")))
            for match in {m.group(0): m.group(1) for m in pattern.finditer(text)}.items():
                replace_in_paragraph(p_el, *match)


def find_paragraph(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"Template is missing the '{text}' marker paragraph")


def build_findings(doc, df):
    """Insert one section per severity in place of the {{FINDINGS}} marker."""
    marker = find_paragraph(doc, "{{FINDINGS}}")
    sub_num_id = add_sub_bullet_numbering(doc)
    fails = df[df["Status"].astype(str).str.lower() == "fail"]

    for severity in SEVERITIES:
        sev_df = fails[fails["Severity"].astype(str).str.lower() == severity]
        marker.insert_paragraph_before(f"{severity.title()} Risk Findings", style="Heading 2")

        if sev_df.empty:
            marker.insert_paragraph_before(f"No {severity} risk findings were identified.")

        for check_title, group in sev_df.groupby("Check Title", dropna=False, sort=True):
            marker.insert_paragraph_before(str(check_title), style="Heading 3")

            for title, column in [
                ("Description", "Status Extended"),
                ("Risk", "Risk"),
                ("Affected Resource IDs", "Resource ID"),
                ("Recommendations", "Recommendation"),
            ]:
                marker.insert_paragraph_before(title, style="Heading 4")
                for val in unique_values(group, column) or ["No data found."]:
                    add_bullet_text(marker, val, sub_num_id)

            # Sections Prowler doesn't provide for a check are left out entirely.
            steps = unique_values(group, "Remediation Steps")
            if steps:
                marker.insert_paragraph_before("Remediation Steps", style="Heading 4")
            for block in steps:
                # Console steps sometimes embed a ```bash``` command block
                for is_code, chunk in split_code_blocks(block):
                    if is_code:
                        marker.insert_paragraph_before(style="K Code Block").add_run(chunk.strip())
                        continue
                    for line in chunk.splitlines():
                        if line.strip():
                            marker.insert_paragraph_before(strip_markdown(line))

            for title in ["CLI", "Terraform", "CloudFormation"]:
                codes = unique_values(group, title)
                if codes:
                    marker.insert_paragraph_before(title, style="Heading 4")
                for code in codes:
                    # "\n" in the run becomes a line break inside the code block
                    marker.insert_paragraph_before(style="K Code Block").add_run(strip_code_fence(code))

        marker.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)

    marker._p.getparent().remove(marker._p)


# --------------------------------------------------------------------------- #
# Native chart refresh
# --------------------------------------------------------------------------- #
def set_cache(ref_el, ref, values, numeric):
    """Rewrite a c:strRef / c:numRef formula and its cached points."""
    ref_el.find("c:f", NSMAP_C).text = ref
    cache = ref_el.find("c:numCache" if numeric else "c:strCache", NSMAP_C)
    for pt in cache.findall("c:pt", NSMAP_C):
        cache.remove(pt)
    cache.find("c:ptCount", NSMAP_C).set("val", str(len(values)))
    for i, v in enumerate(values):
        pt = etree.SubElement(cache, f"{{{C_NS}}}pt", idx=str(i))
        etree.SubElement(pt, f"{{{C_NS}}}v").text = str(v)


def update_chart(chart_part, categories, series):
    """series: list of (name, values). Sheet1 holds categories in column A, one column per series."""
    col = openpyxl.utils.get_column_letter
    last_row = len(categories) + 1

    root = etree.fromstring(chart_part.blob)
    for i, ser in enumerate(root.findall(".//c:ser", NSMAP_C)):
        if i >= len(series):
            ser.getparent().remove(ser)
            continue
        name, values = series[i]
        letter = col(i + 2)
        tx = ser.find("c:tx/c:strRef", NSMAP_C)
        if tx is not None:
            set_cache(tx, f"Sheet1!${letter}$1", [name], numeric=False)
        set_cache(ser.find("c:cat/c:strRef", NSMAP_C), f"Sheet1!$A$2:$A${last_row}", categories, numeric=False)
        set_cache(ser.find("c:val/c:numRef", NSMAP_C), f"Sheet1!${letter}$2:${letter}${last_row}", values, numeric=True)
    chart_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Keep the embedded workbook ("Edit Data" in Word) in sync with the chart
    for rel in chart_part.rels.values():
        if not rel.reltype.endswith("/package"):
            continue
        wb = openpyxl.load_workbook(io.BytesIO(rel.target_part.blob))
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
        ws.delete_rows(2, ws.max_row)
        for c, (name, values) in enumerate(series, start=2):
            ws.cell(row=1, column=c, value=name)
            for r, value in enumerate(values, start=2):
                ws.cell(row=r, column=c, value=value)
        for r, cat in enumerate(categories, start=2):
            ws.cell(row=r, column=1, value=cat)
        for table in ws.tables.values():
            table.ref = f"A1:{col(len(series) + 1)}{last_row}"
            for c, column in enumerate(table.tableColumns, start=1):
                column.name = str(ws.cell(row=1, column=c).value)
        buf = io.BytesIO()
        wb.save(buf)
        rel.target_part._blob = buf.getvalue()


def refresh_charts(doc, fails, by_severity, phase):
    charts = [doc.part.related_parts[el.get(qn("r:id"))] for el in doc.element.body.iter(f"{{{C_NS}}}chart")]
    if len(charts) < 3:
        print("[!] WARNING: template has fewer than 3 charts; charts not updated.", file=sys.stderr)
        return

    findings = dict(zip(by_severity["Severity"].str.lower(), by_severity["Findings"]))
    update_chart(charts[0], [phase], [(s.title(), [int(findings.get(s, 0))]) for s in SEVERITIES])

    svc = fails["Service Name"].value_counts()
    update_chart(charts[1], list(svc.index), [("Failed Count", [int(v) for v in svc.values])])

    reg = fails["Region"].value_counts()
    update_chart(charts[2], list(reg.index), [("Failed Count", [int(v) for v in reg.values])])


# --------------------------------------------------------------------------- #
# Narrative text
# --------------------------------------------------------------------------- #
def service_name(name):
    return SERVICE_NAMES.get(str(name).lower(), str(name))


def join_words(items):
    items = list(items)
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]


def narrative(fails, by_severity):
    findings = dict(zip(by_severity["Severity"].str.lower(), by_severity["Findings"]))
    counts = [
        f"{n} {s} risk {'vulnerability' if n == 1 else 'vulnerabilities'}"
        for s in SEVERITIES for n in [int(findings.get(s, 0))]
    ]
    finding_summary = (
        f"Kernellix uncovered {join_words(counts)}. "
        "Figure 4.1 illustrates the summary of findings by category."
    )

    if fails.empty:
        notable = "During the assessment, Kernellix did not identify any failed security checks in the cloud environment."
        return finding_summary, "No failed checks were identified.", "", notable

    svc = [service_name(s) for s in fails["Service Name"].value_counts().index]
    service_summary = (
        f"{svc[0]} and {svc[1]} have the highest volume of failures." if len(svc) > 1
        else f"All failures are in {svc[0]}."
    ) + " Figure 4.2 illustrates the summary of findings by AWS Services."

    reg = list(fails["Region"].value_counts().index)
    region_summary = (
        f"Most of the risk is concentrated in {reg[0]}, followed by {reg[1]}." if len(reg) > 1
        else f"All of the risk is concentrated in {reg[0]}."
    ) + " Figure 4.3 illustrates the summary of findings by AWS Regions."

    n_crit, n_high = int(findings.get("critical", 0)), int(findings.get("high", 0))
    top = fails[fails["Severity"].str.lower().isin(["critical", "high"])]
    label = "critical and high-risk" if n_crit else "high-risk"
    if top.empty:
        top, label = fails, "medium and low-risk"
    top_services = [service_name(s) for s in top["Service Name"].value_counts().index[:5]]
    notable = (
        f"During the assessment, Kernellix identified {n_crit} critical and {n_high} high-risk "
        f"findings across the cloud environment. Notable {label} findings affect "
        f"{join_words(top_services)}. These findings highlight areas where the environment may be "
        "susceptible to external and internal threats, including unauthorized access, privilege "
        "escalation, data exfiltration, and degraded incident detection and response."
    )
    return finding_summary, service_summary, region_summary, notable


# --------------------------------------------------------------------------- #
# Report
# --------------------------------------------------------------------------- #
def fmt_date(d):
    return f"{d.day:02d} {d:%B %Y}"


def mask_account(account):
    if not account or len(account) < 6:
        return account or "[AWS Account]"
    return account[:2] + "*" * (len(account) - 4) + account[-2:]


def enable_update_fields(doc):
    """Ask Word to refresh the TOC, figure/table lists and captions when the file is opened."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def set_document_properties(doc, project, code):
    """The cover title and the footer are content controls bound to these properties;
    Word re-reads them on open, so they must match the project."""
    old_title = doc.core_properties.title or ""
    new_title = f"{project} AWS Cloud Security Assessment"
    doc.core_properties.title = new_title
    doc.core_properties.subject = "Report"
    for rel in doc.part.package.rels.values():
        if rel.reltype.endswith("/extended-properties"):
            part = rel.target_part
            blob = re.sub(rb"<Company>[^<]*</Company>", f"<Company>{escape(code)}</Company>".encode(), part.blob)
            if old_title:  # "Titles of parts" metadata repeats the title
                blob = blob.replace(escape(old_title).encode(), escape(new_title).encode())
            part._blob = blob


def slugify(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def report_date(args):
    return datetime.strptime(args.date, "%d %B %Y").date() if args.date else date.today()


def report_basename(args):
    """rpt.<code>-<project>-cloud-pentest-prelim.<YYYYMMDD>, shared by the docx and findings xlsx."""
    prefix = slugify(re.sub(r"\d+$", "", args.code)) or "rpt"
    # "EM Cloud" -> "em", so the name reads ...-em-cloud-pentest-... not ...-em-cloud-cloud-pentest-...
    project = re.sub(r"(^|-)cloud$", "", slugify(args.project))
    name = "-".join(x for x in [prefix, project, "cloud-pentest-prelim"] if x)
    return f"rpt.{name}.{report_date(args):%Y%m%d}"


def build_report(df, meta, args, out_path, by_severity):
    template = Path(args.template)
    if not template.is_file():
        raise ValueError(f"Template not found: {template}")
    doc = Document(template)

    fails = df[df["Status"] == "FAIL"]
    scan_date = meta.get("scan_date", date.today())
    masked = mask_account(meta.get("account", ""))

    finding_summary, service_summary, region_summary, notable = narrative(fails, by_severity)

    # No dates given -> the scan date. Only --start -> ends on the scan date. Only --end -> one day.
    end = args.end or fmt_date(scan_date)
    start = args.start or (end if args.end else fmt_date(scan_date))
    if start == end:  # one-day assessment: "on 13 October 2026", not "from X to X"
        period = {f"{word} {{{{START_DATE}}}} to {{{{END_DATE}}}}": f"on {start}" for word in ("from", "between")}
    else:
        period = {"between {{START_DATE}} to {{END_DATE}}": f"between {start} and {end}"}

    mapping = {
        **period,
        "{{PROJECT}}": args.project,
        "{{CLIENT}}": args.client,
        "{{CODE}}": args.code,
        "{{START_DATE}}": start,
        "{{END_DATE}}": end,
        "{{REPORT_DATE}}": fmt_date(report_date(args)),
        "{{ACCOUNT}}": masked,
        "{{USER_ID}}": args.user_id,
        "{{CALLER_ARN}}": args.caller_arn or f"arn:aws:iam::{masked}:user/[user]",
        "{{PROWLER_VERSION}}": meta.get("prowler_version") or "[version]",
        "{{FINDING_SUMMARY}}": finding_summary,
        "{{SERVICE_SUMMARY}}": service_summary,
        "{{REGION_SUMMARY}}": region_summary,
        "{{NOTABLE}}": notable,
    }

    replace_everywhere(doc, mapping)
    collapse_repeated_words(doc)
    build_findings(doc, df)
    refresh_charts(doc, fails, by_severity, "Preliminary")
    enable_update_fields(doc)
    set_document_properties(doc, args.project, args.code)

    doc.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# Driver
# --------------------------------------------------------------------------- #
def find_reports(src):
    """Return a list of report files to process. Native CSV preferred; HTML as fallback."""
    if src.is_file():
        # A HTML file has no remediation code; use the CSV of the same scan if present
        sibling_csv = src.with_suffix(".csv")
        if src.suffix.lower() == ".html" and sibling_csv.is_file():
            print(f"[*] Using {sibling_csv.name} instead of the HTML (it contains remediation code)")
            return [sibling_csv]
        return [src]

    csvs = sorted(src.glob("*.csv"))
    csv_stems = {p.stem for p in csvs}
    # Only use the HTML when there is no CSV for the same scan
    htmls = [p for p in sorted(src.glob("*.html")) if p.stem not in csv_stems]
    return csvs + htmls


def process_report(report, out_dir, args):
    print(f"[*] Parsing {report}")
    if report.suffix.lower() == ".html":
        df, meta = parse_prowler_html(report)
    elif report.suffix.lower() == ".csv":
        df, meta = parse_prowler_csv(report)
    else:
        raise ValueError(f"Unsupported file type: {report.name}")

    if not df[CODE_COLUMNS].apply(lambda c: c.astype(str).str.strip().ne("")).any().any():
        print(
            "[!] WARNING: no remediation code found. Prowler's HTML/OCSF output does not include "
            "CLI/Terraform/CloudFormation; put the Prowler .csv in the input folder.",
            file=sys.stderr,
        )

    other = df[(df["Status"] == "FAIL") & ~df["Severity"].str.lower().isin(SEVERITIES)]
    if not other.empty:
        print(
            f"[!] WARNING: {len(other)} FAIL findings with severity "
            f"{sorted(other['Severity'].unique())} are not in the report (only in the xlsx).",
            file=sys.stderr,
        )

    out_dir.mkdir(parents=True, exist_ok=True)

    base = report_basename(args)
    findings_path = out_dir / f"{base}.xlsx"
    count = build_findings_xlsx(df, findings_path)
    print(f"[+] {findings_path} ({count} findings)")

    xlsx_path = out_dir / "finding-count.xlsx"
    by_service, by_region, by_severity = build_count_xlsx(df, xlsx_path)
    print(f"[+] {xlsx_path} (services: {len(by_service)} | regions: {len(by_region)})")

    docx_path = out_dir / f"{base}.docx"
    build_report(df, meta, args, docx_path, by_severity)
    print(f"[+] {docx_path}")

    print(f"[=] Status: {df['Status'].value_counts().to_dict()}")
    print("[=] Findings (affected resources): " + ", ".join(
        f"{s} {n} ({r})" for s, n, r in by_severity.itertuples(index=False)
    ))


def main():
    parser = argparse.ArgumentParser(
        description="Generate a preliminary DOCX report, a findings XLSX and a finding-count XLSX from Prowler output."
    )
    parser.add_argument("-f", "--folder", required=True,
                        help="Folder containing Prowler output (.csv preferred, .html fallback), or a single file")
    parser.add_argument("-o", "--output", required=True, help="Output folder")
    parser.add_argument("-t", "--template", default=str(DEFAULT_TEMPLATE), help="Report template (.docx)")
    parser.add_argument("--project", default="[Project Name]", help='Project name, e.g. "Yoma Elevator"')
    parser.add_argument("--client", default="Yoma Strategic Holdings Ltd.", help="Client name (Prepared for)")
    parser.add_argument("--code", default="YSH01", help="Client code shown in the header/footer")
    parser.add_argument("--start", help='Assessment start date, e.g. "01 June 2026" (default: same as end date)')
    parser.add_argument("--end", help='Assessment end date, e.g. "04 June 2026" (default: scan date)')
    parser.add_argument("--date", help='Report date, e.g. "29 June 2026" (default: today)')
    parser.add_argument("--user-id", default="[AWS User ID]", help="AWS user ID used for the scan (scope table)")
    parser.add_argument("--caller-arn", help="Caller identity ARN used for the scan (scope table)")
    args = parser.parse_args()

    src = Path(args.folder)
    out = Path(args.output)
    if not src.exists():
        sys.exit(f"Error: input not found: {src}")

    reports = find_reports(src)
    if not reports:
        sys.exit(f"Error: no Prowler .html or .csv report found in {src}")

    for report in reports:
        target = out if len(reports) == 1 else out / report.stem
        try:
            process_report(report, target, args)
        except ValueError as e:
            print(f"[!] Skipping {report.name}: {e}", file=sys.stderr)

    print(f"[✓] Done. Reports written to {out}")


if __name__ == "__main__":
    main()
